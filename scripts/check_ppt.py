"""W22 自检：PPT 的结构与内容。

查什么
------
1. 页数 8–10、文件名符合课程要求（`20_题3_黄锦添.pptx`）；
2. **每页都有一句话结论**（页眉那条），且不超过 34 字（渲染实测 34 字以上会溢出色带）；
3. 每页都有备注（备注就是逐字稿的唯一出处）；
4. 每页都有图或足够的信息量，不许出现「只有标题」的空页；
5. 数字与产物一致（抽查幻灯片正文里的关键数）。

用法
----
    python scripts/check_ppt.py
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu, Inches

ROOT = Path(__file__).resolve().parents[1]
PPTX = ROOT / "report" / "20_题3_黄锦添.pptx"
OUT = ROOT / "outputs"
MAX_MSG = 34


def main() -> int:
    assert PPTX.exists(), f"缺 PPT：{PPTX}"
    prs = Presentation(str(PPTX))
    fails: list[str] = []

    def need(cond: bool, msg: str) -> None:
        if not cond:
            fails.append(msg)

    n = len(prs.slides)
    need(8 <= n <= 10, f"页数 {n} 不在 8–10 之间")
    need(PPTX.name == "20_题3_黄锦添.pptx", f"文件名不符合要求：{PPTX.name}")

    all_text: list[str] = []
    for i, slide in enumerate(prs.slides, 1):
        texts = [sh.text_frame.text.strip() for sh in slide.shapes if sh.has_text_frame]
        texts = [t for t in texts if t]
        pics = [sh for sh in slide.shapes if sh.shape_type == 13]  # PICTURE
        notes = slide.notes_slide.notes_text_frame.text.strip()

        need(notes, f"第 {i} 页没有备注（逐字稿无从生成）")
        need(texts, f"第 {i} 页没有任何文字")
        # 第 1 页是封面，不套「一句话结论」的规矩
        if i > 1:
            # 页眉 = 顶部色带里最长的那块文字（排除小标签与页码）
            band = []
            for sh in slide.shapes:
                if sh.has_text_frame and sh.top is not None and sh.top < Emu(Inches(1.35)):
                    t = sh.text_frame.text.strip()
                    if len(t) > 4:
                        band.append(t)
            msg = max(band, key=len) if band else ""
            need(bool(msg), f"第 {i} 页找不到页眉的一句话结论")
            need(len(msg) <= MAX_MSG,
                 f"第 {i} 页的页眉 {len(msg)} 字，超过 {MAX_MSG} 字会溢出（{msg[:20]}…）")
            need(len(texts) >= 3, f"第 {i} 页只有 {len(texts)} 个文本块，信息量不足")
            need(len(pics) >= 1 or i == 3 or i == 9,
                 f"第 {i} 页既没有图，也不是纯文字页——排版可能塌了")
        all_text.extend(texts)

    body = "\n".join(all_text)

    # 数字抽查：幻灯片里出现的关键数与产物一致
    from pandas import read_csv

    strat = read_csv(OUT / "metrics" / "strategy.csv")
    s1 = strat[(strat["表"] == "S1") & (strat["判据"] == "SUM 裕度") & strat["真值"].notna()]
    lo = s1["提前量比例"].min() * 100
    hi = s1["提前量比例"].max() * 100
    need(f"{lo:.0f}%" in body and f"{hi:.0f}%" in body,
         f"策略页的提前量 {lo:.0f}%–{hi:.0f}% 没出现在幻灯片里（或与产物不一致）")
    for needle in ("0.997", "63%"):
        need(needle in body, f"幻灯片里找不到关键数字 {needle}")

    if fails:
        print()
        for f in fails:
            print("  [FAIL]", f)
        print(f"\n共 {len(fails)} 项不通过")
        return 1

    chars = sum(len(t) for t in all_text)
    print(f"页数 {n}；文件名 {PPTX.name}；每页均有备注与一句话结论")
    print(f"幻灯片文字合计 {chars} 字（不含备注）；数字抽查通过")
    return 0


if __name__ == "__main__":
    sys.exit(main())
