"""W21 自检：报告与产物是否对得上。

为什么要有这个脚本
------------------
报告是唯一一份「散文里带着数字」的交付物，而本项目被审查抓到最多的问题就是
**数字算对了、据它写出的论述是编的**。所以这里做三件事：

1. **结构**：9 章齐全、图 ≥12、表 ≥5、目录是 Word 域；
2. **引用一致**：正文里出现的「表 N / 图 N」必须都有对应的表题/图题；
3. **数字可追**：正文里断言的每个关键数字，都要能在 `outputs/` 的产物里找到来源。
   给不出来源的，就是 `[未核实]`。

用法
----
    python scripts/check_report.py
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

import pandas as pd
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "report" / "20_题3_黄锦添_项目报告.docx"
OUT = ROOT / "outputs"

CHAPTERS = ("1. 引言", "2. 数据介绍", "3. 数据清洗与预处理", "4. 探索性分析与个体差异",
            "5. 模型构建与对比", "6. 不确定性量化、预警与维护策略", "7. 总结与结论",
            "8. 小组成员分工与个人承担", "9. 附录")


def main() -> int:
    assert DOCX.exists(), f"缺报告：{DOCX}（先跑 python scripts/build_report.py）"
    doc = Document(str(DOCX))
    text = "\n".join(p.text for p in doc.paragraphs)
    paras = [p.text.strip() for p in doc.paragraphs]

    # ---- 1. 结构 ----
    h1 = [p.text.strip() for p in doc.paragraphs if p.style.name == "Heading 1"]
    for ch in CHAPTERS:
        assert ch in h1, f"缺章：{ch}（现有 {h1}）"
    fig_caps = [t for t in paras if re.match(r"^图 \d+　", t)]
    tab_caps = [t for t in paras if re.match(r"^表 \d+　", t)]
    assert len(fig_caps) >= 12, f"图题只有 {len(fig_caps)} 个，少于 12"
    assert len(tab_caps) >= 5, f"表题只有 {len(tab_caps)} 个，少于 5"
    assert len(doc.tables) == len(tab_caps), (
        f"表题 {len(tab_caps)} 个但实际表格 {len(doc.tables)} 个——有题无表或有表无题"
    )
    assert len(doc.inline_shapes) == len(fig_caps), (
        f"图题 {len(fig_caps)} 个但实际图片 {len(doc.inline_shapes)} 张"
    )
    assert "TOC" in doc.element.xml, "目录不是 Word 域（F9 无法更新）"

    # ---- 2. 引用一致：正文提到的图表号必须有题 ----
    fig_nums = {int(re.match(r"^图 (\d+)", t).group(1)) for t in fig_caps}
    tab_nums = {int(re.match(r"^表 (\d+)", t).group(1)) for t in tab_caps}
    for m in re.finditer(r"图 (\d+)", text):
        n = int(m.group(1))
        assert n in fig_nums, f"正文引用了图 {n}，但没有这张图"
    for m in re.finditer(r"表 (\d+)", text):
        n = int(m.group(1))
        assert n in tab_nums, f"正文引用了表 {n}，但没有这张表"

    # ---- 3. 数字可追：每个断言都要在「正文」和「产物」里同时出现 ----
    checks: list[tuple[str, str, bool]] = []

    def claim(desc: str, needle: str, source_md: str) -> None:
        """正文里有这个数，并且对应的产物 md 里也有——两头对得上才算可追。"""
        in_doc = needle in text
        src = (OUT / "tables" / source_md).read_text(encoding="utf-8")
        in_src = needle in src
        checks.append((f"{desc}（来源 {source_md}）", needle, in_doc and in_src))

    claim("容量再生最大回升", "0.1519", "regeneration_events.md")
    claim("B0007 最低容量 1.4005", "1.4005", "strategy.md")
    claim("M1 的 EOL 估计", "146", "model_comparison.md")
    claim("随机森林训练 R² 上限", "0.9987", "model_rf.md")
    claim("90% 名义区间实际覆盖", "63%", "interval_calibration.md")
    claim("误差放大上限", "10.9", "error_amplification.md")
    claim("SUM 提前量比例下限", "26.3%", "strategy.md")
    claim("任务 B 的零基准 MAE", "14.7", "task_b_lobocv.md")

    bad = [c for c in checks if not c[2]]
    for desc, needle, ok in checks:
        print(("  [OK]   " if ok else "  [FAIL] ") + f"{desc} -> {needle!r}")
    assert not bad, f"{len(bad)} 条数字断言在正文或产物里找不到"

    print()
    print(f"结构：{len(h1)} 个一级标题（目录 + 9 章）；图 {len(fig_caps)}；表 {len(tab_caps)}")
    print(f"正文 {len(text)} 字；图表引用全部有对应题注；{len(checks)} 条数字可追")
    return 0


if __name__ == "__main__":
    sys.exit(main())
